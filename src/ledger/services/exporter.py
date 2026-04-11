from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ledger.storage.repository import EntryRecord, MatchRecord

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def export_workbook(entries: list[EntryRecord], matches: list[MatchRecord], path: Path) -> None:
    workbook = Workbook()

    entries_sheet = workbook.active
    entries_sheet.title = "分录状态"
    entries_sheet.append(
        [
            "ID",
            "行号",
            "摘要",
            "凭证字号",
            "借方金额",
            "贷方金额",
            "row_type",
            "amount_side",
            "original_amount",
            "matched_amount",
            "remaining_amount",
            "status",
        ]
    )
    _style_header_row(entries_sheet, 1)

    for entry in entries:
        entries_sheet.append(
            [
                entry.id,
                entry.row_index,
                entry.summary,
                entry.voucher_no,
                float(entry.debit_amount),
                float(entry.credit_amount),
                entry.row_type,
                entry.amount_side,
                float(entry.original_amount),
                float(entry.matched_amount),
                float(entry.remaining_amount),
                entry.status,
            ]
        )

    matches_sheet = workbook.create_sheet("匹配记录")
    matches_sheet.append(
        [
            "匹配ID",
            "匹配类型",
            "匹配模式",
            "匹配金额",
            "原因",
            "创建者",
            "创建时间",
            "涉及分录",
        ]
    )
    _style_header_row(matches_sheet, 1)

    for match in matches:
        summary = ", ".join(
            f"{link.role}:{link.entry_id}/{link.used_amount:,.2f}"
            for link in match.links
        )
        matches_sheet.append(
            [
                match.id,
                match.match_type,
                match.match_mode,
                float(match.matched_amount),
                match.reason,
                match.created_by,
                match.created_at,
                summary,
            ]
        )

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _style_header_row(ws, row: int) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
