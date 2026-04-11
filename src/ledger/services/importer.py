from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from ledger.config import REQUIRED_LEDGER_COLUMNS, TARGET_SHEET_NAME


@dataclass(frozen=True)
class ImportedLedgerRow:
    row_index: int
    subject_code: str
    subject_name: str
    posted_date: date | None
    business_date: date | None
    voucher_no: str
    summary: str
    debit_amount: Decimal
    credit_amount: Decimal
    direction: str
    balance: Decimal
    row_type: str
    amount_side: str
    original_amount: Decimal
    remaining_amount: Decimal


@dataclass(frozen=True)
class FileImportError:
    file_path: Path
    code: str
    message: str


@dataclass
class ImportResult:
    rows: list[ImportedLedgerRow]
    file_errors: list[FileImportError]

    def __init__(self) -> None:
        self.rows = []
        self.file_errors = []

    @property
    def records_imported(self) -> int:
        return len(self.rows)


def import_paths(paths: list[Path]) -> ImportResult:
    result = ImportResult()
    for path in [Path(p) for p in paths]:
        if path.is_dir():
            for child in sorted(path.iterdir()):
                _import_supported_path(child, result)
        else:
            _import_supported_path(path, result)
    return result


def _import_supported_path(file_path: Path, result: ImportResult) -> None:
    suffix = file_path.suffix.lower()
    if suffix == ".xlsx":
        _import_xlsx(file_path, result)
    elif suffix == ".csv":
        _import_csv(file_path, result)


def _import_xlsx(file_path: Path, result: ImportResult) -> None:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        if TARGET_SHEET_NAME not in workbook.sheetnames:
            result.file_errors.append(
                FileImportError(
                    file_path=file_path,
                    code="missing_sheet",
                    message=f"找不到工作表「{TARGET_SHEET_NAME}」",
                )
            )
            return

        sheet = workbook[TARGET_SHEET_NAME]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return

        headers = [str(v).strip() if v is not None else "" for v in rows[0]]
        missing = [c for c in REQUIRED_LEDGER_COLUMNS if c not in headers]
        if missing:
            result.file_errors.append(
                FileImportError(
                    file_path=file_path,
                    code="missing_columns",
                    message=f"缺少列：{', '.join(missing)}",
                )
            )
            return

        state = _CarryForwardState()
        for row_index, values in enumerate(rows[1:], start=1):
            raw = dict(zip(headers, values, strict=False))
            normalized = _normalize(raw, row_index=row_index, carry_state=state)
            result.rows.append(normalized)
    finally:
        workbook.close()


def _import_csv(file_path: Path, result: ImportResult) -> None:
    with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        missing = [c for c in REQUIRED_LEDGER_COLUMNS if c not in headers]
        if missing:
            result.file_errors.append(
                FileImportError(
                    file_path=file_path,
                    code="missing_columns",
                    message=f"缺少列：{', '.join(missing)}",
                )
            )
            return

        state = _CarryForwardState()
        for row_index, raw in enumerate(reader, start=1):
            result.rows.append(_normalize(dict(raw), row_index=row_index, carry_state=state))


class _CarryForwardState:
    def __init__(self) -> None:
        self.subject_code = ""
        self.subject_name = ""


_OPENING_KEYWORDS = ("年初余额", "期初余额")
_PERIOD_TOTAL_KEYWORDS = ("本期合计", "本期发生额", "合计")
_YEAR_TOTAL_KEYWORDS = ("本年累计",)


def _normalize(
    raw: dict[str, object],
    *,
    row_index: int,
    carry_state: _CarryForwardState,
) -> ImportedLedgerRow:
    subject_code = _carry_forward_text(raw.get("科目代码"), carry_state.subject_code)
    subject_name = _carry_forward_text(raw.get("科目名称"), carry_state.subject_name)
    carry_state.subject_code = subject_code or carry_state.subject_code
    carry_state.subject_name = subject_name or carry_state.subject_name

    summary = str(raw.get("摘要", "")).strip()
    debit_amount = _parse_decimal(raw.get("借方金额"))
    credit_amount = _parse_decimal(raw.get("贷方金额"))
    row_type = _detect_row_type(summary)
    amount_side = _detect_amount_side(debit_amount, credit_amount)
    original_amount = _calculate_original_amount(row_type, debit_amount, credit_amount)
    remaining_amount = original_amount if row_type == "entry" else Decimal("0.00")

    return ImportedLedgerRow(
        row_index=row_index,
        subject_code=subject_code,
        subject_name=subject_name,
        posted_date=_parse_date(raw.get("日期")),
        business_date=_parse_date(raw.get("业务日期")),
        voucher_no=str(raw.get("凭证字号", "")).replace(" ", ""),
        summary=summary,
        debit_amount=debit_amount,
        credit_amount=credit_amount,
        direction=str(raw.get("方向", "")).strip(),
        balance=_parse_decimal(raw.get("余额")),
        row_type=row_type,
        amount_side=amount_side,
        original_amount=original_amount,
        remaining_amount=remaining_amount,
    )


def _carry_forward_text(value: object, previous: str) -> str:
    text = str(value).strip() if value is not None else ""
    return text or previous


def _detect_row_type(summary: str) -> str:
    normalized = summary.replace(" ", "").replace("\u3000", "")
    if any(keyword in normalized for keyword in _OPENING_KEYWORDS):
        return "opening"
    if any(keyword in normalized for keyword in _YEAR_TOTAL_KEYWORDS):
        return "year_total"
    if any(keyword in normalized for keyword in _PERIOD_TOTAL_KEYWORDS):
        return "period_total"
    return "entry"


def _detect_amount_side(debit_amount: Decimal, credit_amount: Decimal) -> str:
    if debit_amount > 0:
        return "debit"
    if credit_amount > 0:
        return "credit"
    return ""


def _calculate_original_amount(
    row_type: str,
    debit_amount: Decimal,
    credit_amount: Decimal,
) -> Decimal:
    if row_type != "entry":
        return Decimal("0.00")
    if debit_amount > 0:
        return debit_amount
    if credit_amount > 0:
        return credit_amount
    return Decimal("0.00")


def _parse_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        serial = int(float(text))
        return date(1899, 12, 30) + timedelta(days=serial)
    except (ValueError, OverflowError):
        return None


def _parse_decimal(value: object) -> Decimal:
    if value is None:
        return Decimal("0.00")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(round(value, 2)))
    text = str(value).strip().replace(",", "")
    if not text:
        return Decimal("0.00")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0.00")
